# -*- coding: utf-8 -*-
#
# -----------------------------------------------------------------------------
# Nome do Arquivo: Executar.py
#
# Autor:         [Aldair José]
# E-mail:        [aldairsjr@gmail.com]
# Data:          [Data de Criação: 30/04/2026]
# Versão:        [1.8]
# -----------------------------------------------------------------------------

import os
import pandas as pd
from datetime import datetime
import time
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
import subprocess
import re
import requests
import traceback

import undetected_chromedriver as uc
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import StaleElementReferenceException

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics

import pdfplumber

EXCEL_FILE = 'Extratos - Banco do Brasil.xlsx'
SHEET_NAME = 'Extratos'
METRICAS_FILE = 'Métricas de Tempo.xlsx'
URL_BB = 'https://www63.bb.com.br/portalbb/djo/id/resgate/dadosResgate.bbx'
NOME_COLUNA_ADICIONAL = 'Conta Judicial'
PASTA_DOWNLOADS = 'Baixados'
NOME_COLUNA_PRECATORIO = 'Precatorio'

def registrar_metrica(tipo, identificador, tempo_gasto):
    """
    Registra o tempo gasto para busca ou tabulação em um arquivo Excel separado.
    """

    """Definição de estilos"""
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    summary_hdr_fill = PatternFill(start_color="244062", end_color="244062", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_normal = Font(name="Calibri", size=11)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    
    if not os.path.exists(METRICAS_FILE):
        workbook = openpyxl.Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
    else:
        try:
            workbook = openpyxl.load_workbook(METRICAS_FILE)
        except Exception:
            workbook = openpyxl.Workbook()
            default_sheet = workbook.active
            workbook.remove(default_sheet)

    if tipo not in workbook.sheetnames:
        sheet = workbook.create_sheet(title=tipo)
        sheet.views.sheetView[0].showGridLines = True
        
        """Cabeçalhos do Log Detalhado"""
        headers = ["Data/Hora", "Identificador", "Tempo (segundos)"]
        for col_idx, h in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
            
        """Cabeçalhos da Tabela de Resumo Estatístico (Colunas E e F)"""
        sheet.cell(row=1, column=5, value="Métrica Resumo").font = font_header
        sheet.cell(row=1, column=5).fill = summary_hdr_fill
        sheet.cell(row=1, column=5).alignment = left_align
        sheet.cell(row=1, column=5).border = thin_border
        
        sheet.cell(row=1, column=6, value="Valor").font = font_header
        sheet.cell(row=1, column=6).fill = summary_hdr_fill
        sheet.cell(row=1, column=6).alignment = center_align
        sheet.cell(row=1, column=6).border = thin_border
        
        indicadores = ["Total de Quantidade", "Média de Tempo (s)"]
        for idx, ind in enumerate(indicadores, 2):
            c_ind = sheet.cell(row=idx, column=5, value=ind)
            c_ind.font = font_bold
            c_ind.border = thin_border
            c_ind.alignment = left_align
            
            c_val = sheet.cell(row=idx, column=6)
            c_val.font = font_bold
            c_val.border = thin_border
            c_val.alignment = center_align
    else:
        sheet = workbook[tipo]

    """Localizar a próxima linha disponível no log"""
    proxima_linha = 2
    while sheet.cell(row=proxima_linha, column=1).value is not None:
        proxima_linha += 1

    """Gravar as informações do item individual"""
    data_atual = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    valores_log = [data_atual, identificador, round(tempo_gasto, 2)]
    
    for col_idx, val in enumerate(valores_log, 1):
        cell = sheet.cell(row=proxima_linha, column=col_idx, value=val)
        cell.font = font_normal
        cell.alignment = center_align
        cell.border = thin_border
        if proxima_linha % 2 == 0:
            cell.fill = zebra_fill

    """Atualizar as fórmulas do painel de resumo utilizando referências do Excel"""
    sheet.cell(row=2, column=6, value=f"=COUNTA(A2:A{proxima_linha})").number_format = '#,##0'
    sheet.cell(row=3, column=6, value=f"=AVERAGE(C2:C{proxima_linha})").number_format = '0.00'

    """Ajustar as larguras das colunas para garantir legibilidade automática"""
    larguras_colunas = {'A': 22, 'B': 25, 'C': 18, 'D': 4, 'E': 22, 'F': 15}
    for col, largura in larguras_colunas.items():
        sheet.column_dimensions[col].width = largura

    try:
        workbook.save(METRICAS_FILE)
        print(f"Métrica de [{tipo}] salva com sucesso para: {identificador} ({tempo_gasto:.2f}s)")
    except Exception as e:
        print(f"Erro ao salvar arquivo de métricas de tempo: {e}")


def salvar_dados_preservando_formatacao(filename, sheetname, df_final):
    center_align = Alignment(horizontal='center', vertical='center')

    try:
        workbook = openpyxl.load_workbook(filename)
        sheet = workbook[sheetname]

        current_headers = [c.value for c in sheet[1]]

        if NOME_COLUNA_ADICIONAL not in current_headers:
            cell = sheet.cell(row=1, column=sheet.max_column + 1, value=NOME_COLUNA_ADICIONAL)
            cell.alignment = center_align
            print(f"Coluna '{NOME_COLUNA_ADICIONAL}' adicionada à planilha")
            current_headers.append(NOME_COLUNA_ADICIONAL)

        if NOME_COLUNA_PRECATORIO not in current_headers:
            cell = sheet.cell(row=1, column=sheet.max_column + 1, value=NOME_COLUNA_PRECATORIO)
            cell.alignment = center_align
            print(f"Coluna '{NOME_COLUNA_PRECATORIO}' adicionada à planilha")
        
        rows = dataframe_to_rows(df_final, index=False, header=False)

        for r_idx, row_data in enumerate(rows, 2):
            for c_idx, value in enumerate(row_data, 1):
                cell = sheet.cell(row=r_idx, column=c_idx, value=value)
                cell.alignment = center_align
        
        if (sheet.max_row - 1) > len(df_final):
            sheet.delete_rows(len(df_final) + 2, sheet.max_row - (len(df_final)))

        workbook.save(filename)
        print("Progresso salvo")

    except Exception as e:
        print(f"Erro ao salvar com openpyxl: {e}")


def salvar_como_pdf(nome_arquivo, dados_texto):
    """Salva os dados de texto em um arquivo PDF."""
    try:
        try:
            pdfmetrics.registerFont(TTFont('CourierNew', 'cour.ttf'))
            font_name = 'CourierNew'
        except:
            print("Fonte 'Courier New' não encontrada, usando 'Courier' padrão")
            font_name = 'Courier'
        
        caminho_completo = os.path.join(os.getcwd(), PASTA_DOWNLOADS, f"{nome_arquivo}.pdf")
        
        doc = SimpleDocTemplate(caminho_completo, pagesize=letter)
        styles = getSampleStyleSheet()
        style = styles['Normal']
        style.fontName = font_name
        style.fontSize = 9
        style.leading = 11
        style.alignment = TA_LEFT
        story = [Paragraph(linha.replace(' ', '&nbsp;'), style) for linha in dados_texto]
        doc.build(story)
        return f"{nome_arquivo}.pdf"
    except Exception as e:
        print(f"Erro ao salvar Extrato PDF: {e}")
        return None


def obter_versao_chrome():
    """Descobre a versão principal do Chrome instalada no Windows."""
    try:
        comando = r'reg query "HKEY_CURRENT_USER\Software\Google\Chrome\BLBeacon" /v version'
        resultado = subprocess.check_output(comando, shell=True).decode()
        versao = re.search(r'\d+\.\d+\.\d+\.\d+', resultado).group()
        versao_principal = int(versao.split('.')[0])
        return versao_principal
    except Exception:
        return None

def iniciar_driver():
    try:
        options = uc.ChromeOptions()
        options.add_argument('--disable-gpu')
        options.add_argument('--log-level=3')
        
        versao_detectada = obter_versao_chrome()
        
        if versao_detectada:
            print(f"Chrome detectado: Versão {versao_detectada}")
            driver = uc.Chrome(options=options, use_subprocess=True, version_main=versao_detectada)
        else:
            driver = uc.Chrome(options=options, use_subprocess=True)
            
        return driver
    except Exception as e:
        print(f"Erro ao iniciar driver: {e}")
        return uc.Chrome(options=options, use_subprocess=True)


def enviar_dados_para_google_sheets():
    """Lê a planilha local e envia todos os dados para o Google Sheets."""
    print("Enviando dados para a Planilha do Google...")
    URL_WEB_APP = 'https://script.google.com/macros/s/AKfycbxzO4g1dbGmzNBsYt4hk9ZtcrKmfSDwdP37688nOTpfij_LHs2grEbZ4nbBfLqvUZLK/exec'

    if URL_WEB_APP == '':
        print("Erro: A URL do Web App não está definida")
        return

    try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_NAME, dtype=str)
        df.fillna('', inplace=True)
        payload = {
            'headers': list(df.columns),
            'rows': df.to_dict(orient='records')
        }
        response = requests.post(URL_WEB_APP, json=payload)
        response.raise_for_status()
        response_data = response.json()
        if response_data.get('status') == 'sucesso':
            print("Dados enviados com sucesso")
        else:
            print(f"Erro ao enviar dados para a planilha do Google: {response_data.get('message')}")
    except Exception as e:
        print(f"Falha na comunicação com o Google Planilhas: {e}")


def processar_extratos_localmente():
    print("Iniciando processamento...")
    time.sleep(2)
    os.makedirs(PASTA_DOWNLOADS, exist_ok=True)
    
    try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_NAME, dtype=str)
        df.fillna('', inplace=True)
    except FileNotFoundError:
        print(f"Erro: Arquivo '{EXCEL_FILE}' não encontrado")
        return

    if 'CNPJ de Busca' not in df.columns:
        print(f"Erro: A coluna 'CNPJ de Busca' não foi encontrada na planilha '{EXCEL_FILE}'")
        print("Por favor, adicione a coluna e preencha os dados necessários")
        return

    if NOME_COLUNA_ADICIONAL not in df.columns:
        df[NOME_COLUNA_ADICIONAL] = ''
    if NOME_COLUNA_PRECATORIO not in df.columns:
        df[NOME_COLUNA_PRECATORIO] = ''

    linhas_para_processar = df[df.iloc[:, 0] == ''].copy()
    
    if linhas_para_processar.empty:
        print("Nenhum extrato pendente")
        return

    print(f"Total de extratos para processamento: {len(linhas_para_processar)}")
    driver = iniciar_driver()
    
    if not driver:
        print("Processo interrompido por falha ao iniciar o navegador")
        return
    
    try:
        for index, row in linhas_para_processar.iterrows():
            protocolo_da_planilha = row.iloc[2]
            cnpj_da_planilha = row['CNPJ de Busca']

            if not protocolo_da_planilha or not cnpj_da_planilha or pd.isna(protocolo_da_planilha) or pd.isna(cnpj_da_planilha):
                continue

            print(f"Protocolo processado: {protocolo_da_planilha}")
            
            print("Etapa 1 de 2: Baixando extrato...")
            sucesso_download = False
            tentativas = 0
            max_tentativas = 3
            nome_arquivo_pdf = None
            
            """[MÉTRICA] Início do tempo de busca"""
            inicio_busca = time.time()
            
            while not sucesso_download and tentativas < max_tentativas:
                try:
                    driver.get(URL_BB)
                    wait = WebDriverWait(driver, 15)
                    
                    protocolo_8_digitos = str(protocolo_da_planilha)[-8:]
                    
                    """1. Preenche o formulário inicial"""
                    wait.until(EC.element_to_be_clickable((By.ID, 'formulario:tipoPessoa:1'))).click()
                    wait.until(EC.visibility_of_element_located((By.ID, 'formulario:protocolo'))).send_keys(protocolo_8_digitos)
                    wait.until(EC.visibility_of_element_located((By.ID, 'formulario:cnpjPessoa'))).send_keys(cnpj_da_planilha) 
                    wait.until(EC.element_to_be_clickable((By.ID, 'formulario:btnContinuar'))).click()
                    
                    """2. PRIMEIRA TELA DE SELEÇÃO (Mandado / Processo)"""
                    wait.until(EC.presence_of_element_located((By.NAME, 'indice')))
                    wait.until(EC.element_to_be_clickable((By.NAME, 'indice'))).click()
                    wait.until(EC.element_to_be_clickable((By.ID, 'formulario:btnContinuar'))).click()

                    """3. SEGUNDA TELA DE SELEÇÃO (Comparação do Protocolo Resgate)"""
                    wait.until(EC.presence_of_element_located((By.NAME, 'indice')))
                    time.sleep(1)
                    
                    linhas_tabela = driver.find_elements(By.XPATH, "//tr[.//input[@name='indice']]")
                    protocolo_encontrado = False
                    
                    for linha in linhas_tabela:
                        texto_da_linha = linha.text
                        if protocolo_8_digitos in texto_da_linha:
                            botao_radio = linha.find_element(By.XPATH, ".//input[@name='indice']")
                            driver.execute_script("arguments[0].click();", botao_radio)
                            protocolo_encontrado = True
                            break 
                    
                    if not protocolo_encontrado:
                        raise Exception(f"O protocolo {protocolo_8_digitos} não estava visível nas opções da tabela de resgate.")
                        
                    wait.until(EC.element_to_be_clickable((By.ID, 'formulario:btnContinuar'))).click()
                    
                    """4. TELA FINAL (Comprovante e Salvamento do Arquivo)"""
                    wait.until(EC.presence_of_element_located((By.TAG_NAME, 'pre')))
                    
                    elementos_pre = driver.find_elements(By.TAG_NAME, 'pre')
                    dados_comprovante_texto = [elem.text.strip() for elem in elementos_pre]
                    
                    nome_base_arquivo = protocolo_8_digitos.zfill(17)
                    nome_arquivo_pdf = salvar_como_pdf(nome_base_arquivo, dados_comprovante_texto)

                    if nome_arquivo_pdf:
                        df.loc[index, df.columns[0]] = nome_arquivo_pdf
                        print(f"Download concluído: '{nome_arquivo_pdf}'")
                        sucesso_download = True

                except Exception as e:
                    tentativas += 1
                    print(f"Tentativa {tentativas}/{max_tentativas} falhou. Erro: {e}")
                    time.sleep(2)

            """[MÉTRICA] Fim e registro do tempo de busca"""
            fim_busca = time.time()
            if sucesso_download:
                registrar_metrica('Busca', str(protocolo_da_planilha), fim_busca - inicio_busca)

            if not sucesso_download:
                print(f"Falha ao baixar o protocolo {protocolo_da_planilha}. Pulando para o próximo")
                continue

            print("Etapa 2 de 2: Tabulando extrato...")
            caminho_pdf = os.path.join(os.getcwd(), PASTA_DOWNLOADS, nome_arquivo_pdf)
            
            """[MÉTRICA] Início do tempo de tabulação"""
            inicio_tab = time.time()
            
            try:
                with pdfplumber.open(caminho_pdf) as pdf:
                    todas_as_linhas = []
                    for page in pdf.pages:
                        texto_pagina = page.extract_text()
                        if texto_pagina:
                            todas_as_linhas.extend(texto_pagina.split('\n'))

                    colunas_alvo = {col.strip().replace(':', '').lower(): col for col in df.columns[2:]}
                    codigo_adicional_encontrado = False
                    
                    for linha in todas_as_linhas:
                        if not codigo_adicional_encontrado and re.match(r'^\S{13}\s', linha.strip()):
                            df.loc[index, NOME_COLUNA_ADICIONAL] = linha.strip()[:13]
                            codigo_adicional_encontrado = True

                        for chave_norm, nome_col_orig in colunas_alvo.items():
                            if chave_norm in linha.lower():
                                partes = linha.split(':', 1)
                                if len(partes) > 1 and partes[1].strip():
                                    df.loc[index, nome_col_orig] = partes[1].strip()
                
                df.loc[index, df.columns[1]] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                print("Tabulação concluída.")
                
                """[MÉTRICA] Fim e registro do tempo de tabulação"""
                fim_tab = time.time()
                registrar_metrica('Tabulação', str(nome_arquivo_pdf), fim_tab - inicio_tab)

            except Exception as e:
            
                print(f"Falha ao tabular o PDF '{nome_arquivo_pdf}'.")
                print("Detalhes técnicos do erro:")
                traceback.print_exc()
                
                """Registra na planilha que houve um erro para não ficar tentando eternamente"""
                df.loc[index, df.columns[1]] = "ERRO AO LER PDF"
                continue

            salvar_dados_preservando_formatacao(EXCEL_FILE, SHEET_NAME, df)

    finally:
        if driver:
            driver.quit()
        print("Processo finalizado")


def tabular_apenas_extratos_baixados():
    print("Iniciando tabulação...")
    time.sleep(2)
    
    try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_NAME, dtype=str)
        df.fillna('', inplace=True)
    except FileNotFoundError:
        print(f"Erro: Arquivo '{EXCEL_FILE}' não encontrado")
        return

    if NOME_COLUNA_ADICIONAL not in df.columns:
        df[NOME_COLUNA_ADICIONAL] = ''
    if NOME_COLUNA_PRECATORIO not in df.columns:
        df[NOME_COLUNA_PRECATORIO] = ''

    print(f"Verificando arquivos na pasta '{PASTA_DOWNLOADS}'...")
    
    if not os.path.exists(PASTA_DOWNLOADS):
        print(f"A pasta '{PASTA_DOWNLOADS}' não existe. Nenhum arquivo para processar")
        return

    arquivos_na_planilha = set(df.iloc[:, 0].dropna())
    arquivos_na_pasta = [f for f in os.listdir(PASTA_DOWNLOADS) if f.lower().endswith('.pdf')]
    arquivos_novos = [f for f in arquivos_na_pasta if f not in arquivos_na_planilha]

    if arquivos_novos:
        novas_linhas = []
        for arquivo in arquivos_novos:
            nova_linha = {df.columns[0]: arquivo}
            for col in df.columns[1:]:
                nova_linha[col] = ''
            novas_linhas.append(nova_linha)
        
        df_novos = pd.DataFrame(novas_linhas)
        df = pd.concat([df, df_novos], ignore_index=True)
    else:
        print("Nenhum arquivo novo encontrado na pasta para adicionar")

    linhas_para_tabular = df[(df.iloc[:, 0] != '') & (df.iloc[:, 1] == '')].copy()
    
    if linhas_para_tabular.empty:
        print("Nenhum extrato pendente de tabulação")
        if arquivos_novos:
            salvar_dados_preservando_formatacao(EXCEL_FILE, SHEET_NAME, df)
        return

    print(f"Quantidade total de extratos para tabular: {len(linhas_para_tabular)}")
    print("Iniciando processo de tabulação... ")
    time.sleep(3)
    
    for index, row in linhas_para_tabular.iterrows():
        nome_arquivo_pdf = row.iloc[0]
        caminho_pdf = os.path.join(os.getcwd(), PASTA_DOWNLOADS, nome_arquivo_pdf)

        print(f"Tabulando arquivo: {nome_arquivo_pdf}")

        if not os.path.exists(caminho_pdf):
            print(f"Arquivo não encontrado na pasta '{PASTA_DOWNLOADS}'. Pulando...")
            continue
        
        """[MÉTRICA] Início do tempo de tabulação avulsa"""
        inicio_tab = time.time()
        
        try:
            with pdfplumber.open(caminho_pdf) as pdf:
                todas_as_linhas = []
                for page in pdf.pages:
                    texto_pagina = page.extract_text()
                    if texto_pagina:
                        todas_as_linhas.extend(texto_pagina.split('\n'))

                colunas_alvo = {col.strip().replace(':', '').lower(): col for col in df.columns[2:]}
                codigo_adicional_encontrado = False
                
                for linha in todas_as_linhas:
                    if not codigo_adicional_encontrado and re.match(r'^\S{13}\s', linha.strip()):
                        df.loc[index, NOME_COLUNA_ADICIONAL] = linha.strip()[:13]
                        codigo_adicional_encontrado = True

                    for chave_norm, nome_col_orig in colunas_alvo.items():
                        if chave_norm == 'cnpj de busca' and pd.notna(row.get('CNPJ de Busca')) and row.get('CNPJ de Busca') != '':
                            continue
                            
                        if chave_norm in linha.lower():
                            partes = linha.split(':', 1)
                            if len(partes) > 1 and partes[1].strip():
                                df.loc[index, nome_col_orig] = partes[1].strip()
            
            df.loc[index, df.columns[1]] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            print("Tabulação concluída")
            
            """[MÉTRICA] Fim e registro do tempo de tabulação avulsa"""
            fim_tab = time.time()
            registrar_metrica('Tabulação', str(nome_arquivo_pdf), fim_tab - inicio_tab)
            
        except Exception as e:
            
            print(f"Falha ao tabular o PDF '{nome_arquivo_pdf}'.")
            print("Detalhes técnicos do erro:")
            traceback.print_exc()
            
            """Registra na planilha que houve um erro para não ficar tentando eternamente"""
            df.loc[index, df.columns[1]] = "ERRO AO LER PDF"
            continue
            
    print("Salvando dados na planilha...")
    salvar_dados_preservando_formatacao(EXCEL_FILE, SHEET_NAME, df)
    print("Processo de tabulação finalizado")


def renomear_arquivos():
    print("Verificando planilha para renomeação de arquivos...")
    try:
        workbook = openpyxl.load_workbook(EXCEL_FILE)
        if SHEET_NAME not in workbook.sheetnames:
            print(f"Erro: A aba '{SHEET_NAME}' não foi encontrada.")
            return
            
        sheet = workbook[SHEET_NAME]
        header = [cell.value for cell in sheet[1]]
        
        if 'Novo Nome' not in header:
            nova_coluna_idx = sheet.max_column + 1
            cell = sheet.cell(row=1, column=nova_coluna_idx, value='Novo Nome')
            center_align = Alignment(horizontal='center', vertical='center')
            cell.alignment = center_align
            workbook.save(EXCEL_FILE)
            header.append('Novo Nome')
            print("Coluna 'Novo Nome' foi criada na sua planilha. Preencha ela para renomear seus arquivos.")
            return

        idx_nome_antigo = 1
        idx_novo_nome = header.index('Novo Nome') + 1

        arquivos_renomeados = 0
        
        for row_idx in range(2, sheet.max_row + 1):
            nome_antigo = sheet.cell(row=row_idx, column=idx_nome_antigo).value
            novo_nome_base = sheet.cell(row=row_idx, column=idx_novo_nome).value
            
            if novo_nome_base and str(novo_nome_base).strip():
                novo_nome_com_extensao = f"{str(novo_nome_base).strip()}.pdf"
                
                if nome_antigo:
                    caminho_antigo = os.path.join(PASTA_DOWNLOADS, str(nome_antigo))
                    caminho_novo = os.path.join(PASTA_DOWNLOADS, novo_nome_com_extensao)
                    
                    if not os.path.exists(caminho_antigo):
                        print(f"Arquivo ignorado (não encontrado na pasta): {nome_antigo}")
                        continue
                        
                    if os.path.exists(caminho_novo):
                        print(f"Arquivo ignorado (já existe um PDF com esse novo nome): {novo_nome_com_extensao}")
                        continue
                        
                    try:
                        os.rename(caminho_antigo, caminho_novo)
                        print(f"Arquivo '{nome_antigo}' renomeado para '{novo_nome_com_extensao}'")
                        
                        sheet.cell(row=row_idx, column=idx_nome_antigo).value = novo_nome_com_extensao
                        sheet.cell(row=row_idx, column=idx_novo_nome).value = None
                        arquivos_renomeados += 1
                        
                    except Exception as e:
                        print(f"Erro ao renomear '{nome_antigo}': {e}")
                        
        if arquivos_renomeados > 0:
            workbook.save(EXCEL_FILE)
            print(f"Planilha salva. {arquivos_renomeados} arquivo(s) renomeado(s).")
        else:
            print("Nenhum arquivo para renomear encontrado na planilha.")
            print("Dica: Preencha a coluna 'Novo Nome' com os nomes desejados antes de executar.")
            
    except Exception as e:
        print(f"Erro ao tentar renomear arquivos a partir da planilha: {e}")
        
    print("Processo de renomeação concluído")

"""BLOCO DE EXECUÇÃO"""

if __name__ == "__main__":
    while True:
        print("Extratos - Banco do Brasil")
        print("")
        print("1 - Buscar e tabular extratos")
        print("2 - Tabular extratos")
        print("3 - Renomear extratos")
        print("4 - Enviar extratos para o Google Planilhas")
        print("5 - Execução completa (Buscar, tabular e enviar)")
        print("6 - Sair")
        
        escolha = input("Digite a opção desejada: ")
        
        if escolha == '1':
            processar_extratos_localmente()
        elif escolha == '2':
            tabular_apenas_extratos_baixados()
        elif escolha == '3':
            renomear_arquivos()
        elif escolha == '4':
            enviar_dados_para_google_sheets()
        elif escolha == '5':
            processar_extratos_localmente()
            enviar_dados_para_google_sheets()
        elif escolha == '6':
            print("Encerrando o programa...")
            break
        else:
            print("Opção inválida. Digite uma das opções listadas!")
